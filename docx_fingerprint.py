import re
import json
from collections import Counter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

PADRAO_TITULO = re.compile(r"(t[íi]tulo|heading)\s*([1-3])\b", re.IGNORECASE)
ROTULOS_CAIXA = (
    ("ATENCAO", "atencao"),
    ("DICA", "dica"),
    ("JURISPRUD", "jurisprudencia"),
    ("EXEMPLO", "exemplo"),
)

ROTULOS_ICONE = (
    ("O PULO DO GATO", "o_pulo_do_gato"),
    ("PEGADINHA DA BANCA", "pegadinha_da_banca"),
    ("PEGADINHA", "pegadinha_da_banca"),
    ("DIRETO DO CONCURSO", "direto_do_concurso"),
    ("EXERCICIOS DE FIXACAO", "exercicios_de_fixacao"),
    ("EXERC", "exercicios_de_fixacao"),
    ("COMENTARIO", "comentario"),
    ("COMENT", "comentario"),
    ("RESOLUCAO", "resolucao"),
    ("RESOLU", "resolucao"),
    ("RELEMBRANDO", "relembrando"),
    ("GABARITO", "gabarito"),
    ("ATENCAO", "atencao"),
    ("OBS", "obs"),
)

PADRAO_RELOGIO = re.compile(r"\b\d{1,2}\s*min(?:utos)?\b", re.IGNORECASE)


def normalizar_rotulo(texto):
    u = re.sub(r"\s+", " ", (texto or "").strip()).upper()
    for chave, sub in (
        ("Ç", "C"), ("Ã", "A"), ("Õ", "O"), ("Ô", "O"), ("Ó", "O"),
        ("É", "E"), ("Ê", "E"), ("Í", "I"), ("Á", "A"), ("Ú", "U"),
    ):
        u = u.replace(chave, sub)
    return u


def rotulo_icone(texto):
    u = normalizar_rotulo(texto)
    for chave, simbolo in ROTULOS_ICONE:
        if u.startswith(chave):
            return simbolo
    return None


def nulo(obj, atributo):
    try:
        valor = getattr(obj, atributo)
        if valor is None:
            return None
        if hasattr(valor, "pt"):
            return round(valor.pt, 2)
        return valor
    except Exception:
        return None


def norm(valor):
    if isinstance(valor, float):
        return round(valor, 2)
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, dict):
        return {k: norm(v) for k, v in valor.items()}
    if isinstance(valor, (list, tuple)):
        return [norm(v) for v in valor]
    return valor


def modo(valores):
    valores = [v for v in valores if v is not None]
    if not valores:
        return None
    return Counter(valores).most_common(1)[0][0]


def resumir_formato(dados):
    return {
        "fonte": modo([x["fonte"] for x in dados["formato_run"]]),
        "tamanho_pt": modo([x["tamanho_pt"] for x in dados["formato_run"]]),
        "negrito": modo([x["negrito"] for x in dados["formato_run"]]),
        "italico": modo([x["italico"] for x in dados["formato_run"]]),
        "sublinhado": modo([x["sublinhado"] for x in dados["formato_run"]]),
        "maiusculas": modo([x["maiusculas"] for x in dados["formato_run"]]),
        "cor": modo([x["cor_hex"] for x in dados["formato_run"]]),
        "subscrito": modo([x["subscrito"] for x in dados["formato_run"]]),
        "sobrescrito": modo([x["sobrescrito"] for x in dados["formato_run"]]),
        "alinhamento": modo([x["alinhamento"] for x in dados["formato_par"]]),
        "espaco_antes_pt": modo([x["espaco_antes_pt"] for x in dados["formato_par"]]),
        "espaco_depois_pt": modo([x["espaco_depois_pt"] for x in dados["formato_par"]]),
        "entrelinha": modo([x["entrelinha"] for x in dados["formato_par"]]),
        "entrelinha_regra": modo([x["entrelinha_regra"] for x in dados["formato_par"]]),
        "recuo_esquerda_pt": modo([x["recuo_esquerda_pt"] for x in dados["formato_par"]]),
        "recuo_primeira_linha_pt": modo([x["recuo_primeira_linha_pt"] for x in dados["formato_par"]]),
        "manter_com_proximo": modo([x["manter_com_proximo"] for x in dados["formato_par"]]),
    }


def cor_rgb(cor):
    try:
        return "#%02X%02X%02X" % tuple(cor)
    except Exception:
        return None


def cor_run(run):
    c = run.font.color
    if c is None or c.type is None:
        return None
    try:
        if c.rgb is not None:
            return cor_rgb(c.rgb)
    except Exception:
        pass
    try:
        if c.theme_color is not None:
            return "tema:%s" % c.theme_color
    except Exception:
        pass
    return None


def formato_run(run, estilo):
    fonte = run.font.name
    if not fonte and estilo is not None and estilo.font.name:
        fonte = estilo.font.name
    tamanho = None
    if run.font.size is not None:
        tamanho = round(run.font.size.pt, 2)
    elif estilo is not None and estilo.font.size is not None:
        tamanho = round(estilo.font.size.pt, 2)
    sub = None
    sup = None
    try:
        sub = run.font.subscript
        sup = run.font.superscript
    except Exception:
        pass
    return {
        "fonte": fonte,
        "tamanho_pt": tamanho,
        "negrito": bool(run.bold),
        "italico": bool(run.italic),
        "sublinhado": bool(run.underline),
        "cor_hex": cor_run(run),
        "maiusculas": bool(run.font.all_caps),
        "cor_alto_contraste": run.font.highlight_color.name if run.font.highlight_color is not None else None,
        "subscrito": bool(sub),
        "sobrescrito": bool(sup),
    }


def formato_paragrafo(p):
    pf = p.paragraph_format
    entrelinha = None
    regra_linha = None
    try:
        if pf.line_spacing_rule is not None:
            regra_linha = pf.line_spacing_rule.name
        ls = pf.line_spacing
        if isinstance(ls, (int, float)):
            entrelinha = round(float(ls), 2)
        elif ls is not None:
            entrelinha = round(ls.pt, 2)
    except Exception:
        pass
    return {
        "alinhamento": pf.alignment.name if pf.alignment is not None else None,
        "espaco_antes_pt": nulo(pf, "space_before"),
        "espaco_depois_pt": nulo(pf, "space_after"),
        "entrelinha": entrelinha,
        "entrelinha_regra": regra_linha,
        "recuo_esquerda_pt": nulo(pf, "left_indent"),
        "recuo_direita_pt": nulo(pf, "right_indent"),
        "recuo_primeira_linha_pt": nulo(pf, "first_line_indent"),
        "manter_com_proximo": bool(pf.keep_with_next),
    }


def sombra_paragrafo(p):
    pPr = p._p.pPr
    if pPr is None:
        return None
    shd = pPr.find(qn("w:shd"))
    if shd is None:
        return None
    return {
        "preenchimento": shd.get(qn("w:fill")),
        "valor": shd.get(qn("w:val")),
        "cor": shd.get(qn("w:color")),
    }


def borda_paragrafo(p):
    pPr = p._p.pPr
    if pPr is None:
        return None
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        return None
    out = {}
    for lado in ("top", "bottom", "left", "right"):
        el = pBdr.find(qn("w:%s" % lado))
        if el is not None:
            sz = el.get(qn("w:sz"))
            out[lado] = {
                "estilo": el.get(qn("w:val")),
                "cor": el.get(qn("w:color")),
                "espaco": el.get(qn("w:space")),
                "largura_pt": (int(sz) / 8) if sz else None,
            }
    return out


def sombra_celula(cel):
    tcPr = cel._tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return None
    return {
        "preenchimento": shd.get(qn("w:fill")),
        "valor": shd.get(qn("w:val")),
        "cor": shd.get(qn("w:color")),
    }


def nivel_titulo(p):
    m = PADRAO_TITULO.search(p.style.name or "")
    if m:
        return int(m.group(2))
    return None


def eh_corpo(p):
    nome = (p.style.name or "").lower()
    return bool(re.match(r"^(normal|corpo|body|list)", nome))


def eh_titulo_principal(p, texto, negrito_total):
    if not texto or len(texto) < 4 or not negrito_total:
        return False
    try:
        return p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        return False


def eh_enfase(p, texto):
    texto = texto or ""
    if len(texto) < 3:
        return False
    runs = [r for r in p.runs if r.text.strip()]
    if not runs:
        return False
    if all(bool(r.bold) for r in runs) and not all(bool(r.italic) for r in runs):
        return not p.text.isdigit() and not texto[:1].isdigit()
    return False


def eh_citacao(texto):
    t = re.sub(r"[\s\u00a0]+", " ", (texto or "").strip()).upper()
    return bool(
        t.startswith(("ART.", "ART ", "ARTIGO", "§", "SUUMLULA", "SÚMULA", "LEI ", "DECRETO", "CF/88", "INCISO"))
        or re.match(r"^\s*\(§", t)
    )


def rotulo_tipo(texto):
    u = re.sub(r"\s+", " ", (texto or "").strip()).upper()
    for chave, simbolo in ROTULOS_CAIXA:
        if u.startswith(chave):
            return simbolo
    return None


def rotulo_corpo(texto):
    u = re.sub(r"\s+", " ", (texto or "").strip()).upper()
    if u.startswith(("MINHA DICA", "DICA")):
        return "dica"
    return None


def contagem_elementos(raiz, extensao):
    return sum(1 for el in raiz.iter() if el.tag.rsplit("}", 1)[-1] == extensao)


def estilos_nativos(doc):
    out = {}
    for est in doc.styles:
        m = PADRAO_TITULO.search(est.name or "")
        if not m:
            continue
        try:
            cor = None
            if est.font.color is not None and est.font.color.type is not None:
                try:
                    if est.font.color.rgb is not None:
                        cor = cor_rgb(est.font.color.rgb)
                except Exception:
                    pass
                try:
                    if est.font.color.theme_color is not None:
                        cor = "tema:%s" % est.font.color.theme_color
                except Exception:
                    pass
            pf = est.paragraph_format
            ls = None
            try:
                ls = pf.line_spacing
                if isinstance(ls, (int, float)):
                    ls = round(float(ls), 2)
                elif ls is not None:
                    ls = round(ls.pt, 2)
            except Exception:
                pass
            out["titulo_%s" % m.group(2)] = {
                "fonte": est.font.name,
                "tamanho_pt": round(est.font.size.pt, 2) if est.font.size else None,
                "negrito": bool(est.font.bold),
                "italico": bool(est.font.italic),
                "cor": cor,
                "alinhamento": pf.alignment.name if pf.alignment is not None else None,
                "antes_pt": nulo(pf, "space_before"),
                "depois_pt": nulo(pf, "space_after"),
                "entrelinha": ls,
            }
        except Exception:
            continue
    return out


def analisar_caixa(itens_agrupados):
    rotulo = None
    bordas = {}
    fundos = []
    alinhamentos = []
    fonte_rotulo = []
    tam_rotulo = []
    negrito_rotulo = []
    cor_rotulo = []
    maiusculas_rotulo = []
    fonte_corpo = []
    tam_corpo = []
    negrito_corpo = []
    italico_corpo = []

    primeira_rotulo = True
    for item in itens_agrupados:
        if item.get("borda"):
            for lado, det in item["borda"].items():
                bordas.setdefault(lado, det)
        if item.get("sombra") and item["sombra"]["preenchimento"] not in (None, "auto"):
            fundos.append(item["sombra"]["preenchimento"])
        if item.get("rotulo"):
            rotulo = item["rotulo"]
            for r in item["paragrafo"].runs:
                if not r.text.strip():
                    continue
                fonte_rotulo.append(r.font.name)
                tam_rotulo.append(round(r.font.size.pt, 2) if r.font.size else None)
                negrito_rotulo.append(bool(r.bold))
                cor_rotulo.append(cor_run(r))
                maiusculas_rotulo.append(bool(r.font.all_caps))
        if not item.get("rotulo"):
            for r in item["paragrafo"].runs:
                if not r.text.strip():
                    continue
                fonte_corpo.append(r.font.name)
                tam_corpo.append(round(r.font.size.pt, 2) if r.font.size else None)
                negrito_corpo.append(bool(r.bold))
                italico_corpo.append(bool(r.italic))
        alinhamentos.append(item["paragrafo"].paragraph_format.alignment.name if item["paragrafo"].paragraph_format.alignment else None)

    resumo_borda = {}
    for lado, det in bordas.items():
        resumo_borda[lado] = det
    return {
        "rotulo": rotulo,
        "bordas": resumo_borda,
        "fundo_cor": modo(fundos),
        "formato_rotulo": {
            "fonte": modo(fonte_rotulo),
            "tamanho_pt": modo(tam_rotulo),
            "negrito": modo(negrito_rotulo),
            "cor": modo(cor_rotulo),
            "maiusculas": modo(maiusculas_rotulo),
        },
        "formato_corpo": {
            "fonte": modo(fonte_corpo),
            "tamanho_pt": modo(tam_corpo),
            "negrito": modo(negrito_corpo),
            "italico": modo(italico_corpo),
        },
        "alinhamento": modo(alinhamentos),
    }


def extrair_fingerprint(caminho):
    doc = Document(caminho)

    secoes = []
    cont_desenhos = {"pagina": contagem_elementos(doc.element, "drawing"), "cabecalho": 0, "rodape": 0}
    for s in doc.sections:
        ls = None
        try:
            ls = s.page_width
            ls = round(ls.pt, 2)
        except Exception:
            ls = None
        ap = None
        try:
            ap = s.page_height
            ap = round(ap.pt, 2)
        except Exception:
            ap = None
        m = {}
        for nome, el in (("superior", s.top_margin), ("inferior", s.bottom_margin), ("esquerda", s.left_margin), ("direita", s.right_margin)):
            try:
                m[nome] = round(el.pt, 2)
            except Exception:
                m[nome] = None
        txt_cab = " ".join(p.text for p in s.header.paragraphs).strip()
        txt_rod = " ".join(p.text for p in s.footer.paragraphs).strip()
        cont_desenhos["cabecalho"] += contagem_elementos(s.header._element, "drawing")
        cont_desenhos["rodape"] += contagem_elementos(s.footer._element, "drawing")
        secoes.append(
            {
                "orientacao": "paisagem" if (ls or 0) > (ap or 0) else "retrato",
                "largura_pt": ls,
                "altura_pt": ap,
                "margens_pt": m,
                "cabecalho_texto": txt_cab[:200] or None,
                "rodape_texto": txt_rod[:200] or None,
            }
        )

    paragrafos = list(doc.paragraphs)
    categorias = {}
    amostras_categoria = {}
    caixa_aberta = None
    caixas = []
    citacoes = []
    icones = {}
    bordas_achadas = []
    relogios = []
    total_num = 0

    tamanhos_corpo = []
    for p in paragrafos:
        if eh_corpo(p):
            for r in p.runs:
                if r.text.strip() and r.font.size is not None:
                    tamanhos_corpo.append(round(r.font.size.pt, 2))
    tamanho_corpo_ref = modo(tamanhos_corpo)

    for p in paragrafos:
        texto = p.text.strip()
        borda = borda_paragrafo(p)
        sombra = sombra_paragrafo(p)
        rotulo = rotulo_tipo(texto)
        if borda:
            bordas_achadas.append(borda)
        if PADRAO_RELOGIO.search(texto):
            relogios.append(texto[:120])
        tem_marcador = bool(borda or sombra or rotulo)
        pPr = p._p.pPr
        tem_num = pPr is not None and pPr.find(qn("w:numPr")) is not None
        if tem_num:
            total_num += 1

        negrito_total = bool(p.runs) and all(bool(r.bold) for r in p.runs if r.text.strip())
        icone = rotulo_icone(texto)
        if icone and (negrito_total or len(texto) < 60 or icone == "obs"):
            ic = icones.setdefault(icone, {"formato_run": [], "formato_par": []})
            for r in p.runs:
                if r.text.strip():
                    ic["formato_run"].append(formato_run(r, p.style))
            ic["formato_par"].append(formato_paragrafo(p))

        if tem_marcador:
            if caixa_aberta is None:
                caixa_aberta = []
            item = {"paragrafo": p, "borda": borda, "sombra": sombra, "rotulo": rotulo}
            caixa_aberta.append(item)
            continue
        if caixa_aberta is not None:
            caixas.append(analisar_caixa(caixa_aberta))
            caixa_aberta = None

        estilo = p.style
        nivel = nivel_titulo(p)
        tams = [r.font.size.pt for r in p.runs if r.text.strip() and r.font.size]
        negrito_total = bool(p.runs) and all(bool(r.bold) for r in p.runs if r.text.strip())
        if nivel is not None:
            categoria = "titulo_%s" % nivel
        elif eh_titulo_principal(p, texto, negrito_total):
            categoria = "titulo_principal"
        elif eh_enfase(p, texto):
            categoria = "enfase"
        elif eh_corpo(p):
            categoria = "corpo"
        else:
            categoria = "outro"

        if categoria in ("titulo_principal", "titulo_1", "titulo_2", "titulo_3", "enfase", "corpo"):
            dados = categorias.setdefault(categoria, {"formato_run": [], "formato_par": []})
            for r in p.runs:
                if r.text.strip():
                    dados["formato_run"].append(formato_run(r, estilo))
            dados["formato_par"].append(formato_paragrafo(p))
            amostras_categoria.setdefault(categoria, []).append(texto[:160])

        if eh_citacao(texto) or (texto.startswith(("Art", "art")) and (formato_paragrafo(p)["recuo_esquerda_pt"] or False)):
            citacoes.append(
                {"texto": texto[:160], "paragrafo": formato_paragrafo(p), "runs": [formato_run(r, estilo) for r in p.runs if r.text.strip()]}
            )

    if caixa_aberta is not None:
        caixas.append(analisar_caixa(caixa_aberta))

    caixas_rotuladas = {}
    for cx in caixas:
        chave = cx["rotulo"] or "sem_rotulo"
        caixas_rotuladas.setdefault(chave, []).append(cx)

    caixas_final = {}
    for chave, lista in caixas_rotuladas.items():
        caixas_final[chave] = {
            "quantidade": len(lista),
            "bordas": lista[-1]["bordas"],
            "fundo_cor": modo([x["fundo_cor"] for x in lista]),
            "formato_rotulo": {campo: modo([x["formato_rotulo"].get(campo) for x in lista]) for campo in ("fonte", "tamanho_pt", "negrito", "cor", "maiusculas")},
            "formato_corpo": {campo: modo([x["formato_corpo"].get(campo) for x in lista]) for campo in ("fonte", "tamanho_pt", "negrito", "italico")},
            "alinhamento": modo([x["alinhamento"] for x in lista]),
        }

    tabelas = []
    for tb in doc.tables:
        tblPr = tb._tbl.tblPr
        tblBorders = tblPr.find(qn("w:tblBorders")) if tblPr is not None else None
        bordas = {}
        if tblBorders is not None:
            for lado in ("top", "bottom", "left", "right", "insideH", "insideV"):
                el = tblBorders.find(qn("w:%s" % lado))
                if el is not None:
                    sz = el.get(qn("w:sz"))
                    bordas[lado] = {"estilo": el.get(qn("w:val")), "cor": el.get(qn("w:color")), "largura_pt": (int(sz) / 8) if sz else None}
        cabecalho = None
        rotulos_celula = {}
        if tb.rows:
            prim = tb.rows[0]
            if prim.cells:
                c0 = prim.cells[0]
                sombra_c = sombra_celula(c0)
                neg_rot = all(bool(r.bold) for r in c0.paragraphs[0].runs) if c0.paragraphs else None
                cabecalho = {
                    "texto": c0.text.strip()[:120] or None,
                    "sombra": sombra_c,
                    "negrito": neg_rot,
                    "formato": [formato_run(r, c0.paragraphs[0].style) for r in c0.paragraphs[0].runs] if c0.paragraphs else [],
                }
            for row in tb.rows:
                for cel in row.cells:
                    rt = rotulo_tipo(cel.text)
                    if rt and rt in ("atencao", "dica", "jurisprudencia", "exemplo"):
                        rotulos_celula.setdefault(rt, {"sombra": sombra_celula(cel), "texto": cel.text.strip()[:120]}).setdefault(
                            "quantidade", 0
                        )
                        rotulos_celula[rt]["quantidade"] += 1
        if rotulos_celula:
            for rt, det in rotulos_celula.items():
                if rt in caixas_final:
                    caixas_final[rt]["quantidade"] += det.get("quantidade", 0)
                    caixas_final[rt].setdefault("celula_tabela", False)
                    caixas_final[rt]["celula_tabela"] = True
                    if caixas_final[rt]["fundo_cor"] is None:
                        caixas_final[rt]["fundo_cor"] = det["sombra"]["preenchimento"] if det["sombra"] else None
        tabelas.append({"linhas": len(tb.rows), "colunas": len(tb.columns), "bordas": bordas, "cabecalho": cabecalho})

    categorias_final = {}
    for cat, dados in categorias.items():
        resumo = resumir_formato(dados)
        resumo["amostras"] = amostras_categoria.get(cat, [])[:5]
        categorias_final[cat] = resumo

    icones_final = {}
    for icone, dados in icones.items():
        resumo = resumir_formato(dados)
        resumo["quantidade"] = len(dados["formato_par"])
        icones_final[icone] = resumo

    bottoms = [b.get("bottom") for b in bordas_achadas if b.get("bottom")]
    linha_limitadora = {
        "quantidade": len(bottoms),
        "estilo": modo([b["estilo"] for b in bottoms]),
        "cor": modo([b["cor"] for b in bottoms]),
        "largura_pt": modo([b["largura_pt"] for b in bottoms]),
    }
    marcadores_tempo = {"quantidade": len(relogios), "exemplos": relogios[:5]}

    citacao_resumo = {}
    if citacoes:
        cor_cit = [c["runs"][0]["cor_hex"] for c in citacoes if c["runs"]]
        fonte_cit = [c["runs"][0]["fonte"] for c in citacoes if c["runs"]]
        tam_cit = [c["runs"][0]["tamanho_pt"] for c in citacoes if c["runs"]]
        ital_cit = [c["runs"][0]["italico"] for c in citacoes if c["runs"]]
        citacao_resumo = {
            "quantidade": len(citacoes),
            "fonte": modo(fonte_cit),
            "tamanho_pt": modo(tam_cit),
            "italico": modo(ital_cit),
            "cor": modo(cor_cit),
            "recuo_esquerda_pt": modo([c["paragrafo"]["recuo_esquerda_pt"] for c in citacoes]),
            "recuo_primeira_linha_pt": modo([c["paragrafo"]["recuo_primeira_linha_pt"] for c in citacoes]),
            "exemplos": [c["texto"] for c in citacoes][:5],
        }

    return norm(
        {
            "secoes": secoes,
            "desenhos": cont_desenhos,
            "estilos_nativos": estilos_nativos(doc),
            "categorias_paragrafo": categorias_final,
            "caixas_destaque": caixas_final,
            "icones": icones_final,
            "linha_limitadora": linha_limitadora,
            "marcadores_tempo": marcadores_tempo,
            "citacoes_lei": citacao_resumo,
            "tabelas": {
                "quantidade": len(tabelas),
                "exemplos": tabelas[:3],
            },
            "num_paragrafos": {"total": len(paragrafos), "listas_ornamentais": total_num},
        }
    )


def _prop_ignoravel(caminho, modo):
    """Propriedades que não devem reprovar quando as aulas comparadas são diferentes."""
    if modo != "aula_diferente":
        return False
    base = caminho.split(".")[-1]
    if base in ("amostras", "exemplos", "cabecalho_texto", "rodape_texto"):
        return True
    if "quantidade" in base:
        return True
    if caminho.startswith("num_paragrafos") or base in ("total", "listas_ornamentais"):
        return True
    if caminho == "desenhos.pagina":
        return True
    return False


def _normalizar_identidade(texto):
    return re.sub(r"[^a-z0-9]+", " ", (texto or "").lower()).strip()


def alinhar_para_aula_diferente(a, b):
    """Remove de ambos os fingerprints os blocos cuja presença é conteúdo da aula
    (ícones, caixas e citações que existem só em um dos materiais), preservando o formato."""
    a = json.loads(json.dumps(a, ensure_ascii=False, default=str))
    b = json.loads(json.dumps(b, ensure_ascii=False, default=str))

    for chave in ("icones", "caixas_destaque"):
        comuns = set(a.get(chave, {}) or {}) & set(b.get(chave, {}) or {})
        if chave in a:
            a[chave] = {k: v for k, v in a[chave].items() if k in comuns}
        if chave in b:
            b[chave] = {k: v for k, v in b[chave].items() if k in comuns}

    def tem_conteudo(fp):
        cx = fp.get("citacoes_lei") or {}
        return bool(cx.get("quantidade")) or bool(cx.get("exemplos"))
    if not (tem_conteudo(a) and tem_conteudo(b)):
        a["citacoes_lei"] = {}
        b["citacoes_lei"] = {}
    return a, b


def sao_mesma_aula(fp_a, fp_b):
    """Detecta se dois fingerprints são da mesma aula pelo TÍTULO PRINCIPAL
    (o cabeçalho da seção é a marca/sumário do curso e não identifica a aula)."""
    def titulos(fp):
        vals = []
        for am in (fp.get("categorias_paragrafo", {}).get("titulo_principal", {}) or {}).get("amostras", []) or []:
            v = _normalizar_identidade(am)
            if len(v) >= 10:
                vals.append(v)
        return vals

    ca, cb = titulos(fp_a), titulos(fp_b)
    for x in ca:
        for y in cb:
            if x == y or (len(x) > 15 and (x in y or y in x)):
                return True
    return False


def comparar_fingerprints(a, b, prefixo="", modo="mesma_aula"):
    linhas = []
    chaves = set(a) | set(b)
    for k in sorted(chaves):
        caminho = "%s.%s" % (prefixo, k) if prefixo else k
        va, vb = norm(a.get(k)), norm(b.get(k))
        if isinstance(va, dict) or isinstance(vb, dict):
            linhas += comparar_fingerprints(va or {}, vb or {}, caminho, modo)
            continue
        if isinstance(va, list) or isinstance(vb, list):
            if isinstance(va, list) and isinstance(vb, list) and (va or vb) and all(isinstance(x, dict) for x in va + vb):
                for idx in range(max(len(va), len(vb))):
                    elem_a = va[idx] if idx < len(va) else {}
                    elem_b = vb[idx] if idx < len(vb) else {}
                    linhas += comparar_fingerprints(elem_a, elem_b, "%s.%d" % (caminho, idx), modo)
                continue
            iguais = (va or []) == (vb or [])
            status = "OK" if iguais else ("OK" if _prop_ignoravel(caminho, modo) else "DIVERGENTE")
            linhas.append({"propriedade": caminho, "amostra": va, "gerado": vb, "status": status,
                           "ignorada": (not iguais and status == "OK")})
            continue
        iguais = va == vb
        status = "OK" if iguais else ("OK" if _prop_ignoravel(caminho, modo) else "DIVERGENTE")
        linhas.append({"propriedade": caminho, "amostra": va, "gerado": vb, "status": status,
                       "ignorada": (not iguais and status == "OK")})
    return linhas