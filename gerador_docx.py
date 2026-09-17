import json
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def adicionar_borda_inferior(paragrafo):
    pPr = paragrafo._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')  # 1.5 pt
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)
    pPr.append(pBdr)

def adicionar_destaque_amarelo(run):
    rPr = run._r.get_or_add_rPr()
    highlight = OxmlElement('w:highlight')
    highlight.set(qn('w:val'), 'yellow')
    rPr.append(highlight)

def _run_base(paragrafo, texto="", negrito=False, italico=False, tamanho=11):
    run = paragrafo.add_run(texto)
    run.font.name = "Arial"
    run.font.size = Pt(tamanho)
    run.bold = negrito
    run.italic = italico
    return run

PADRAO_POTENCIA = re.compile(r"([\^_])(\{[^}]*\}|[A-Za-z0-9]+)")

def _carregar_recuos():
    """Recuos configuráveis por aula via regras_estilo.json (fallback para o padrão)."""
    recuos = {
        "corpo_esquerda_pt": 35.4,
        "corpo_primeira_linha_pt": 35.4,
        "citacao_esquerda_pt": 72.0,
    }
    caminho = os.path.join(os.path.dirname(os.path.abspath(__file__)), "regras_estilo.json")
    try:
        with open(caminho, "r", encoding="utf-8") as f:
            estilos = json.load(f).get("estilos", {})
        corpo = estilos.get("corpo") or {}
        if corpo.get("recuo_esquerda_pt") is not None:
            recuos["corpo_esquerda_pt"] = float(corpo["recuo_esquerda_pt"])
        if corpo.get("recuo_primeira_linha_pt") is not None:
            recuos["corpo_primeira_linha_pt"] = float(corpo["recuo_primeira_linha_pt"])
    except Exception:
        pass
    return recuos

def adicionar_formula(paragrafo, texto, tamanho=11):
    """Adiciona o texto interpretando ^ como sobrescrito e _ como subscrito."""
    pos = 0
    for m in PADRAO_POTENCIA.finditer(texto):
        if m.start() > pos:
            _run_base(paragrafo, texto[pos:m.start()], tamanho=tamanho)
        alvo = m.group(2)
        if alvo.startswith("{"):
            alvo = alvo[1:-1]
        run = _run_base(paragrafo, alvo, tamanho=tamanho)
        if m.group(1) == "^":
            run.font.superscript = True
        else:
            run.font.subscript = True
        pos = m.end()
    if pos < len(texto):
        _run_base(paragrafo, texto[pos:], tamanho=tamanho)

def adicionar_tabela(doc, bloco):
    cabecalho = bloco.get("cabecalho") or []
    linhas = bloco.get("linhas") or []
    n_cols = len(cabecalho) if cabecalho else max((len(l) for l in linhas), default=0)
    n_linhas = len(linhas) + (1 if cabecalho else 0)
    if n_cols == 0 or n_linhas == 0:
        return
    tabela = doc.add_table(rows=n_linhas, cols=n_cols)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        tabela.style = "Table Grid"
    except Exception:
        pass
    linha_inicial = 0
    if cabecalho:
        for j, valor in enumerate(cabecalho[:n_cols]):
            celula = tabela.rows[0].cells[j]
            celula.text = ""
            _run_base(celula.paragraphs[0], str(valor), negrito=True)
        linha_inicial = 1
    for i, linha in enumerate(linhas):
        for j in range(n_cols):
            valor = str(linha[j]) if j < len(linha) else ""
            celula = tabela.rows[i + linha_inicial].cells[j]
            celula.text = ""
            _run_base(celula.paragraphs[0], valor)

def gerar_docx(template_path, json_estruturado_path, saida_docx_path, blocos=None):
    print(f"[GeradorDocx] Usando template: {template_path}")
    recuos = _carregar_recuos()
    doc = Document(template_path)
    
    # Limpa o corpo mantendo seções/cabeçalhos/rodapés do template
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith('p') or child.tag.endswith('tbl'):
            body.remove(child)
            
    if blocos is None:
        with open(json_estruturado_path, "r", encoding="utf-8") as f:
            blocos = json.load(f)
        
    print(f"[GeradorDocx] Escrevendo {len(blocos)} blocos...")

    # Linha limitadora (borda inferior) no último parágrafo de conteúdo de cada bloco de ícone
    tipos_conteudo = ("corpo", "questao", "comentario", "gabarito_comentario", "alternativa", "formula", "citacao_lei")
    indices_borda = set()
    for k, bloco in enumerate(blocos):
        if bloco.get("tipo") != "icone":
            continue
        j = k + 1
        while j < len(blocos) and blocos[j].get("tipo") in tipos_conteudo:
            j += 1
        if j - 1 != k:
            indices_borda.add(j - 1)

    for idx, bloco in enumerate(blocos):
        tipo = bloco.get("tipo")
        texto = bloco.get("texto", "")
        
        if tipo == "titulo":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(texto.upper())
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.bold = True
            
        elif tipo == "corpo":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Pt(recuos["corpo_esquerda_pt"])
            p.paragraph_format.first_line_indent = Pt(recuos["corpo_primeira_linha_pt"])
            run = p.add_run(texto)
            run.font.name = "Arial"
            run.font.size = Pt(11)
            if idx in indices_borda:
                adicionar_borda_inferior(p)
            
        elif tipo == "enfase":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(texto)
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.bold = True

        elif tipo == "citacao_lei":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Pt(recuos["citacao_esquerda_pt"])
            _run_base(p, texto, tamanho=10)
            if idx in indices_borda:
                adicionar_borda_inferior(p)
            
        elif tipo == "icone":
            nome = bloco.get("nome", "").replace("_", " ").title()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(nome)
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.bold = True
            
        elif tipo in ("questao", "comentario", "gabarito_comentario"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(texto)
            run.font.name = "Arial"
            run.font.size = Pt(11)
            if idx in indices_borda:
                adicionar_borda_inferior(p)
            
        elif tipo == "alternativa":
            letra = bloco.get("letra", "a")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            rotulo = f"{letra}) {texto}" if letra else texto
            run = p.add_run(rotulo)
            run.font.name = "Arial"
            run.font.size = Pt(11)
            if idx in indices_borda:
                adicionar_borda_inferior(p)

        elif tipo == "formula":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.5
            adicionar_formula(p, texto)
            if idx in indices_borda:
                adicionar_borda_inferior(p)

        elif tipo == "tabela":
            adicionar_tabela(doc, bloco)

        elif tipo == "imagem":
            arquivo = bloco.get("arquivo")
            if not arquivo or not os.path.exists(arquivo):
                print(f"[GeradorDocx] AVISO: imagem não encontrada, ignorando: {arquivo}")
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            try:
                run.add_picture(arquivo, width=Inches(5.5))
            except Exception as erro:
                print(f"[GeradorDocx] AVISO: falha ao inserir imagem {arquivo}: {erro}")
                continue
            legenda = bloco.get("legenda")
            if legenda:
                p_leg = doc.add_paragraph()
                p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run_base(p_leg, legenda, italico=True, tamanho=10)
            
        elif tipo == "tempo":
            minuto = bloco.get("minuto", 5)
            if doc.paragraphs:
                p_last = doc.paragraphs[-1]
                run = p_last.add_run(f" {minuto}MIN")
                run.font.name = "Arial"
                run.font.size = Pt(11)
                run.bold = True
                adicionar_destaque_amarelo(run)
                
        elif tipo == "gabarito_fim":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run("Gabarito")
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.bold = True
            
            respostas = bloco.get("respostas", [])
            for r in respostas:
                p_r = doc.add_paragraph()
                run_r = p_r.add_run(str(r))
                run_r.font.name = "Arial"
                run_r.font.size = Pt(11)

        else:
            print(f"[GeradorDocx] AVISO: bloco de tipo desconhecido ignorado: {tipo!r}")

    # Rodapé padrão do template ou disclaimer final
    p_disc = doc.add_paragraph()
    p_disc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_d = p_disc.add_run("_____________________________________________________________________\nEste material foi elaborado pela equipe pedagógica do Gran Concursos, de acordo com a aula preparada e ministrada.\nA presente degravação tem como objetivo auxiliar no acompanhamento e na revisão do conteúdo ministrado na videoaula.\n_____________________________________________________________________")
    run_d.font.name = "Arial"
    run_d.font.size = Pt(9)

    doc.save(saida_docx_path)
    print(f"[GeradorDocx] Documento final gerado com sucesso em: {saida_docx_path}")

if __name__ == "__main__":
    if len(sys.argv) > 3:
        gerar_docx(sys.argv[1], sys.argv[2], sys.argv[3])
    else:
        print("Uso: python gerador_docx.py <template.docx> <conteudo.json> <saida.docx>")